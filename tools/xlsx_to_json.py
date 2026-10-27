#!/usr/bin/env python3
"""Converte SpaceX_aggiornato.xlsx (file di Master Dati) in data/spacex.json.

Uso:
    python3 tools/xlsx_to_json.py [percorso.xlsx] [--out data/spacex.json]

- Richiede openpyxl. Legge i valori calcolati (data_only=True): il file va
  salvato da Excel/LibreOffice, così le formule hanno il valore in cache.
  Se una formula risulta vuota, ricalcolare con:
      soffice --headless --convert-to xlsx --outdir /tmp/recalc SpaceX_aggiornato.xlsx
- Non modifica mai l'xlsx e non inventa numeri: copia solo i valori presenti.
- I riepiloghi vengono trovati per titolo (es. "Pad di lancio 2026"), non per
  coordinate fisse, così il layout può spostarsi senza rompere lo script.
- Date: nel file sono già in ora di Roma; vengono scritte come "AAAA-MM-GGTHH:MM"
  senza fuso (la pagina le mostra così come sono).
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Serve openpyxl: pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
MESI = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno", "Luglio",
        "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]
WARN = []


def warn(msg):
    WARN.append(msg)
    print("ATTENZIONE:", msg, file=sys.stderr)


# ---------------------------------------------------------------- valori
def clean(v):
    """Valore cella -> JSON (date ISO senza fuso, timedelta in secondi)."""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%dT%H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, dt.timedelta):
        return int(round(v.total_seconds()))
    if isinstance(v, float):
        return int(v) if v.is_integer() else round(v, 4)
    if isinstance(v, str):
        s = v.strip()
        return s if s != "" else None
    return v


def num(v):
    v = clean(v)
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def txt(v):
    v = clean(v)
    return None if v is None else str(v)


def secs(v):
    """Intervalli: timedelta/time -> secondi."""
    v = clean(v)
    return v if isinstance(v, int) else None


def dmy(s):
    """'21/11/2020' o '3/06/2021' -> '2020-11-21'."""
    m = re.match(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$", s or "")
    if not m:
        return None
    return "%04d-%02d-%02d" % (int(m.group(3)), int(m.group(2)), int(m.group(1)))


# ---------------------------------------------------------------- ricerca blocchi
def find_title(ws, title, exact=True, after_col=1):
    for row in ws.iter_rows():
        for c in row:
            if c.column < after_col or not isinstance(c.value, str):
                continue
            v = c.value.strip()
            if (v == title) if exact else v.startswith(title):
                return c.row, c.column
    return None


def block(ws, title, keys=None, exact=True, after_col=1, max_cols=12, stop_blank=1):
    """Legge un riepilogo: titolo, riga intestazioni sotto, righe fino a una riga vuota.
    keys = nomi da dare alle colonne (altrimenti le intestazioni del foglio)."""
    pos = find_title(ws, title, exact, after_col)
    if not pos:
        warn("blocco non trovato: %s / %s" % (ws.title, title))
        return None, []
    r0, c0 = pos
    hdr = []
    for c in range(c0, c0 + max_cols):
        h = ws.cell(r0 + 1, c).value
        if h is None and c > c0:
            break
        hdr.append(txt(h) or "")
    names = keys or hdr
    rows, r, blanks = [], r0 + 2, 0
    while r <= ws.max_row:
        first = ws.cell(r, c0).value
        if first is None or (isinstance(first, str) and not first.strip()):
            blanks += 1
            if blanks >= stop_blank:
                break
            r += 1
            continue
        blanks = 0
        rows.append({names[i]: clean(ws.cell(r, c0 + i).value) for i in range(min(len(names), len(hdr)))})
        r += 1
    return hdr, rows


def kv(rows, k="k", v="v"):
    return [{"label": str(x[k]), "valore": x[v]} for x in rows if x.get(k) is not None]


# ---------------------------------------------------------------- lanci annuali
def missioni(ws, first_row=3):
    """Tabella A–I (#, Missione, Intervallo, Partenza, Booster, Tipologia, Pad, Atterraggio, Vettore).
    Le sotto-righe Falcon Heavy '↳ …' diventano 'altriBooster' della riga principale."""
    out = []
    for r in range(first_row, ws.max_row + 1):
        nome = txt(ws.cell(r, 2).value)
        if not nome:
            continue
        if nome.startswith("↳"):
            if out:
                out[-1].setdefault("altriBooster", []).append(
                    {"booster": txt(ws.cell(r, 5).value), "atterraggio": txt(ws.cell(r, 8).value)})
            continue
        part = clean(ws.cell(r, 4).value)
        if not part:
            continue
        out.append({
            "n": num(ws.cell(r, 1).value),
            "missione": nome,
            "intervallo": secs(ws.cell(r, 3).value),
            "partenza": part,
            "booster": txt(ws.cell(r, 5).value),
            "tipologia": txt(ws.cell(r, 6).value),
            "pad": txt(ws.cell(r, 7).value),
            "atterraggio": txt(ws.cell(r, 8).value),
            "vettore": txt(ws.cell(r, 9).value),
        })
    return out


PAD_TOT = {"Totale", "Falliti", "Missioni operative"}
ATT_TOT = {"Non recuperati", "Fallimenti", "Mare", "Terra", "Tot recuperi", "Tot recuperi F9",
           "Tot tentativi", "Tot atterraggi", "Totale"}


def split_tot(rows, totset):
    voci = [x for x in rows if x["label"] not in totset]
    tot = {x["label"]: x["valore"] for x in rows if x["label"] in totset}
    return voci, tot


def riepiloghi_anno(ws, anno):
    R = {}
    _, rows = block(ws, "Pad di lancio %d" % anno, ["k", "v"])
    R["pad"], R["padTotali"] = split_tot(kv(rows), PAD_TOT)
    _, rows = block(ws, "Piattaforme atterraggio %d" % anno, ["k", "v"])
    R["atterraggi"], R["atterraggiTotali"] = split_tot(kv(rows), ATT_TOT)
    _, rows = block(ws, "Distribuzione missioni %d" % anno, ["k", "v"])
    R["tipologie"] = [x for x in kv(rows) if x["label"].upper() != "TOTALE"]
    R["tipologieTotale"] = next((x["valore"] for x in kv(rows) if x["label"].upper() == "TOTALE"), None)
    _, rows = block(ws, "Lanci %d per mese" % anno, ["k", "v"])
    R["perMese"] = [x for x in kv(rows) if x["label"] in MESI]
    _, rows = block(ws, "Record", ["record", "valore", "booster"])
    R["record"] = [{"record": x["record"], "valore": x["valore"] if not isinstance(x["valore"], int) else fmt_secs(x["valore"]),
                    "booster": x.get("booster")} for x in rows if x.get("record")]
    _, rows = block(ws, "N satelliti Rideshare", ["missione", "satelliti"])
    R["rideshare"] = rows
    return R


def fmt_secs(s):
    d, r = divmod(int(s), 86400)
    h, r = divmod(r, 3600)
    return "%dg %dh %dm" % (d, h, r // 60)


def vettori(ws):
    hdr, rows = block(ws, "Vettori")
    if not rows:
        return None
    nomi = hdr[1:]
    out = []
    for x in rows:
        anno = x.get(hdr[0])
        if not isinstance(anno, int):
            continue
        out.append({"anno": anno, **{n: x.get(n) for n in nomi}})
    return {"vettori": nomi, "anni": out}


def f9_utilizzati(ws):
    hdr, rows = block(ws, "Falcon 9 utilizzati e attivi", max_cols=14)
    out = []
    for x in rows:
        b = {"booster": x.get("Booster"), "perAnno": {}, "totale": x.get("Totale voli"), "note": x.get("Note")}
        for h in hdr:
            m = re.match(r"N voli (\d{4})", h)
            if m and x.get(h) is not None:
                b["perAnno"][m.group(1)] = x.get(h)
        out.append(b)
    return out


# ---------------------------------------------------------------- archivio 2021–2023
def archivio_lanci(wb):
    ws = wb["Lanci"]
    A = {}
    # (anno, col #, col missione, col intervallo, col partenza, col booster, col tipologia, col orbita)
    for anno, c, extra in ((2021, 1, None), (2022, 8, None), (2023, 15, (21, 22, 23))):
        rows = []
        for r in range(3, ws.max_row + 1):
            nome, part = txt(ws.cell(r, c + 1).value), clean(ws.cell(r, c + 3).value)
            if not nome or not part:
                continue
            x = {"n": num(ws.cell(r, c).value), "missione": nome,
                 "intervallo": secs(ws.cell(r, c + 2).value), "partenza": part}
            if extra:
                x["booster"] = txt(ws.cell(r, extra[0]).value)
                x["tipologia"] = txt(ws.cell(r, extra[1]).value)
                x["orbita"] = txt(ws.cell(r, extra[2]).value)
            rows.append(x)
        A[anno] = rows
    # Pad e atterraggi 2022–2023
    wp = wb["Pad & atterragi"]
    for anno, c in ((2022, 1), (2023, 7)):
        pads, byn, last = {}, {}, None
        for r in range(3, wp.max_row + 1):
            nome = txt(wp.cell(r, c + 1).value)
            att = txt(wp.cell(r, c + 4).value)
            if nome:
                last = {"pad": txt(wp.cell(r, c + 3).value), "atterraggio": att, "altri": [], "nome": nome,
                        "partenza": clean(wp.cell(r, c + 2).value)}
                pads[nome] = last
                if num(wp.cell(r, c).value) is not None:
                    byn[num(wp.cell(r, c).value)] = last
            elif att and last:
                last["altri"].append(att)
        for x in A[anno]:
            p = pads.get(x["missione"])
            if not p and x.get("n") in byn and byn[x["n"]]["partenza"] == x["partenza"]:
                # nome scritto diversamente: abbina per numero lancio + stessa data/ora
                p = byn[x["n"]]
                warn("%d #%s: nome diverso tra fogli (stessa data): '%s' (Lanci) / '%s' (Pad & atterragi)" % (anno, x["n"], x["missione"], p["nome"]))
            if p:
                x["pad"], x["atterraggio"] = p["pad"], p["atterraggio"]
                if p["altri"]:
                    x["altriAtterraggi"] = p["altri"]
        senza = [x["missione"] for x in A[anno] if "pad" not in x]
        if senza:
            warn("%d: %d missioni senza pad in 'Pad & atterragi': %s" % (anno, len(senza), ", ".join(senza[:5])))
    return A, wp


def riepiloghi_pad_archivio(wp, anno):
    R = {}
    _, rows = block(wp, "Pad di lancio %d" % anno, ["k", "v"])
    R["pad"], R["padTotali"] = split_tot(kv(rows), PAD_TOT)
    _, rows = block(wp, "Piattaforme atterraggio %d" % anno, ["k", "v"])
    R["atterraggi"], R["atterraggiTotali"] = split_tot(kv(rows), ATT_TOT)
    return R


def riepiloghi_lanci_sheet(wb):
    ws = wb["Lanci"]
    R = {}
    _, rows = block(ws, "Distribuzione missioni 2023", ["k", "v"])
    R["tipologie2023"] = [x for x in kv(rows) if x["label"].upper() != "TOTALE"]
    _, rows = block(ws, "Lanci 2021: Starlink / enti terzi", ["k", "v"])
    R["tipologieTitolo2021"] = [x for x in kv(rows) if x["label"].upper() != "TOTALE"]
    R["tipologieTitolo2021Totale"] = next((x["valore"] for x in kv(rows) if x["label"].upper() == "TOTALE"), None)
    for anno in (2022, 2023):
        _, rows = block(ws, "Lanci %d per mese" % anno, ["k", "v"])
        R["perMese%d" % anno] = [x for x in kv(rows) if x["label"] in MESI]
    _, rows = block(ws, "Lanci per anno", ["k", "v"])
    R["perAnno"] = kv(rows)
    return R


# ---------------------------------------------------------------- Starlink
def starlink(wb):
    ws = wb["Starlink"]
    S = {}
    hdr, rows = block(ws, "Starlink lanciati per mese (n. satelliti)")
    anni = [h for h in hdr[1:] if re.fullmatch(r"\d{4}", h)]
    S["satelliti"] = {"anni": [int(a) for a in anni],
                      "perMese": [{"mese": x["Mese"], **{a: x.get(a) for a in anni}} for x in rows if x["Mese"] in MESI]}
    tot = next((x for x in rows if str(x["Mese"]).upper() == "TOTALE"), None)
    S["satelliti"]["perAnno"] = [{"anno": int(a), "valore": tot.get(a)} for a in anni] if tot else []
    S["satelliti"]["totale"] = tot.get("Totale") if tot else None
    for x in rows:
        if str(x["Mese"]).upper() == "TOTALE GEN1":
            S["satelliti"]["totaleGen1"] = x.get("Totale")
        if str(x["Mese"]).upper() == "TOTALE GEN2":
            S["satelliti"]["totaleGen2"] = x.get("Totale")
    hdr, rows = block(ws, "Missioni Starlink")
    anniM = [h for h in hdr[1:] if re.fullmatch(r"\d{4}", h)]
    S["missioni"] = {"anni": sorted(int(a) for a in anniM),
                     "perMese": [{"mese": x["Mese"], **{a: x.get(a) for a in anniM}} for x in rows if x["Mese"] in MESI]}
    tot = next((x for x in rows if str(x["Mese"]).upper() == "TOTALE"), None)
    S["missioni"]["perAnno"] = sorted(({"anno": int(a), "valore": tot.get(a)} for a in anniM), key=lambda y: y["anno"]) if tot else []
    _, rows = block(ws, "Gusci Orbitali", ["guscio", "inclinazione", "altitudine", "previsti", "lanciati", "ultimoLancio"])
    S["gusci"] = rows
    _, rows = block(ws, "Abbonati", ["data", "abbonati"])
    S["abbonati"] = [x for x in rows if x.get("abbonati") is not None]
    liste = {}
    for anno in (2026, 2025, 2024, 2023):
        hdr, rows = block(ws, "Starlink lanciati per missione %d" % anno, max_cols=6)
        L = []
        for x in rows:
            y = {"missione": x.get("Missione"), "satelliti": x.get("Starlink lanciati"), "massaKg": x.get("Massa (kg)")}
            if x.get("Mese") is not None:
                y["mese"] = x.get("Mese")
            if x.get("Direct to Cell") is not None:
                y["dtc"] = x.get("Direct to Cell")
            L.append(y)
        liste[str(anno)] = L
        dtc = [x.get("TOT DtC") for x in rows if x.get("TOT DtC") is not None]
        if dtc:
            S.setdefault("totDtc", {})[str(anno)] = dtc[0]
    S["liste"] = liste
    return S


# ---------------------------------------------------------------- Falcon 9
def falcon9(wb):
    ws = wb["Falcon 9"]

    def matrice(title_row_label, col_from, col_to):
        res = []
        for c in range(col_from, col_to + 1):
            nome = txt(ws.cell(3, c).value)
            if not nome:
                continue
            voli = []
            for r in range(4, ws.max_row + 1):
                v = txt(ws.cell(r, c).value)
                if not v:
                    continue
                n = num(ws.cell(r, col_from - 1).value)
                m = re.match(r"^(.*?):\s*(\d{1,2}/\d{1,2}/\d{4})\s*$", v)
                voli.append({"n": n, "missione": (m.group(1).strip() if m else v), "data": dmy(m.group(2)) if m else None})
            code = re.search(r"B\d{4}", nome)
            res.append({"nome": nome, "booster": code.group(0) if code else nome,
                        "fh": "FH" in nome, "voli": voli, "nVoli": len(voli)})
        return res
    pos_att = find_title(ws, "Booster attivi")
    pos_persi = find_title(ws, "Booster distrutti")
    end_att = (pos_persi[1] - 1) if pos_persi else ws.max_column
    attivi = matrice("Booster attivi", pos_att[1] + 1, end_att) if pos_att else []
    persi = matrice("Booster distrutti", pos_persi[1] + 1, ws.max_column) if pos_persi else []
    return {"attivi": attivi, "persi": persi}


# ---------------------------------------------------------------- Dragon
def dragon(wb):
    ws = wb["Dragon"]
    caps = []
    for r in range(1, ws.max_row + 1):
        t = txt(ws.cell(r, 1).value)
        if t and ws.cell(r + 1, 1).value == "Missione":
            hdr = [txt(ws.cell(r + 1, c).value) for c in range(1, 8)]
            mis = []
            rr = r + 2
            while rr <= ws.max_row and ws.cell(rr, 1).value:
                x = {"missione": txt(ws.cell(rr, 1).value), "partenza": clean(ws.cell(rr, 2).value),
                     "rientro": clean(ws.cell(rr, 3).value), "durata": txt(ws.cell(rr, 4).value),
                     "turnaround": txt(ws.cell(rr, 5).value)}
                if "Carico Up (kg)" in hdr:
                    x["caricoKg"] = num(ws.cell(rr, hdr.index("Carico Up (kg)") + 1).value)
                mis.append(x)
                rr += 1
            m = re.match(r"^(.*?)\s*\((C\d+)\)$", t)
            nome, codice = (m.group(1), m.group(2)) if m else (t, None)
            cargo = nome.lower().startswith("cargo")
            if cargo:
                codice = nome.split()[-1]
            caps.append({"nome": nome, "codice": codice, "tipo": "cargo" if cargo else "equipaggio", "missioni": mis})
    _, ut = block(ws, "Utilizzi Dragon", ["nome", "voli"])
    _, eq = block(ws, "Equipaggi Dragon", ["n", "missione", "comandante", "pilota", "spec1", "spec2", "capsula"], max_cols=7)
    return {"capsule": caps, "utilizzi": ut, "equipaggi": eq}


# ---------------------------------------------------------------- Kuiper, Modifiche
def kuiper(wb):
    ws = wb["Kuiper"]
    _, rows = block(ws, "Lanci Kuiper", ["missione", "partenza", "vettore", "satelliti", "totale", "tempoStessoVettore", "mediaStessoVettore", "tempoTrascorso"], max_cols=8)
    tot = next((x["totale"] for x in rows if x.get("totale") is not None), None)
    for x in rows:
        x.pop("totale", None)
    return {"lanci": rows, "totaleSatelliti": tot}


def modifiche(wb):
    ws = wb["Modifiche"]
    _, rows = block(ws, "Registro modifiche", ["data", "foglio", "descrizione"])
    for x in rows:
        if isinstance(x["data"], str) and dmy(x["data"]):
            x["data"] = dmy(x["data"])
    return rows


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default="/workspace/SpaceX_aggiornato.xlsx")
    ap.add_argument("--out", default=os.path.join(ROOT, "data", "spacex.json"))
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.xlsx, data_only=True, read_only=False)

    lanci, riep = {}, {}
    for anno in (2026, 2025, 2024):
        ws = wb["Lanci %d" % anno]
        lanci[str(anno)] = missioni(ws)
        riep[str(anno)] = riepiloghi_anno(ws, anno)
    arch, wp = archivio_lanci(wb)
    for anno, rows in arch.items():
        lanci[str(anno)] = rows
    for anno in (2022, 2023):
        riep[str(anno)] = riepiloghi_pad_archivio(wp, anno)
    extra = riepiloghi_lanci_sheet(wb)
    riep["2023"]["tipologie"] = extra["tipologie2023"]
    riep["2023"]["perMese"] = extra["perMese2023"]
    riep["2022"]["perMese"] = extra["perMese2022"]
    riep["2021"] = {}
    # il riepilogo intitolato "Lanci 2021: Starlink / enti terzi" ha totale 61 = lanci 2022:
    # lo esportiamo con il suo titolo originale, senza attribuirlo a un anno.
    archivio_extra = {"riepilogoStarlinkEntiTerzi": {"titoloNelFile": "Lanci 2021: Starlink / enti terzi",
                                                      "voci": extra["tipologieTitolo2021"],
                                                      "totale": extra["tipologieTitolo2021Totale"]},
                      "lanciPerAnnoFoglioLanci": extra["perAnno"]}

    # controllo coerenza: righe tabella vs totale riepilogo
    for anno in ("2026", "2025", "2024"):
        t = riep[anno]["padTotali"].get("Totale")
        if t is not None and t != len(lanci[anno]):
            warn("Lanci %s: %d righe ma totale pad = %s" % (anno, len(lanci[anno]), t))

    mod = modifiche(wb)
    date_mod = sorted(x["data"] for x in mod if isinstance(x.get("data"), str))
    out = {
        "meta": {
            "fonte": os.path.basename(a.xlsx),
            "aggiornato": date_mod[-1] if date_mod else None,
            "generato": dt.datetime.now().strftime("%Y-%m-%dT%H:%M"),
            "avvisi": WARN,
        },
        "lanci": lanci,
        "riepiloghi": riep,
        "vettori": vettori(wb["Lanci 2026"]),
        "f9Utilizzati2026": f9_utilizzati(wb["Lanci 2026"]),
        "falcon9": falcon9(wb),
        "starlink": starlink(wb),
        "dragon": dragon(wb),
        "kuiper": kuiper(wb),
        "archivio": archivio_extra,
        "modifiche": mod,
    }
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print("Scritto %s (%d KB). Lanci: %s. Avvisi: %d" % (
        a.out, os.path.getsize(a.out) // 1024, ", ".join("%s=%d" % (k, len(v)) for k, v in sorted(lanci.items())), len(WARN)))


if __name__ == "__main__":
    main()
